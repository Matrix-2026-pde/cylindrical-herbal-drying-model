# 生成 result1.xlsx（对齐附件3模板：温度/水分浓度两表，A列时间1..1800s，第1行距离0..2cm，四位小数）
import os, sys, json
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)
import numpy as np
from openpyxl import Workbook
from params import OUTPUT_R_CM, OUTPUT_T

RESULT_PATH = 'result1.xlsx'


def write():
    d = json.load(open('figures/problem_1_results.json', encoding='utf-8'))
    T = np.array(d['T_field'])   # [1800 x 21]
    C = np.array(d['C_field'])
    r_cm = OUTPUT_R_CM
    t_s = OUTPUT_T

    wb = Workbook()
    for name, field in (('温度', T), ('水分浓度', C)):
        ws = wb.active if name == '温度' else wb.create_sheet(name)
        ws.title = name
        # 第1行：表头 + 距离
        ws.cell(row=1, column=1, value='时间\\到药材中心的距离')
        for j, rc in enumerate(r_cm):
            ws.cell(row=1, column=2 + j, value=round(float(rc), 4))
        # 数据行：A列时间, 四位小数
        for i, tt in enumerate(t_s):
            ws.cell(row=2 + i, column=1, value=int(tt))
            for j in range(len(r_cm)):
                ws.cell(row=2 + i, column=2 + j, value=round(float(field[i, j]), 4))
    wb.save(RESULT_PATH)
    print("[write_result] 已写 %s  温度/水分浓度 各 %d行×%d列(四位小数)" % (RESULT_PATH, len(t_s), len(r_cm)))


if __name__ == '__main__':
    write()

