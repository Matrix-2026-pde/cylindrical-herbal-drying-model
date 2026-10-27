# -*- coding: utf-8 -*-
# 按附件3模板格式填充 result1-4.xlsx（覆盖到 附件3 目录）
# result1: 温度/水分浓度, t=1..1800 s; result2: 温度/水分浓度, t=1..10800 s (dt=1重算)
# result3: Sheet1 水分浓度, t=60..t* 每60 s; result4: Sheet1 + 药材表面列, t=60..t*_shrink 每60 s
import os, sys, json
import numpy as np
from openpyxl import Workbook

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

OUTDIR = os.path.join(_HERE, '..', '..', '03_结果文件')
R_CM = np.round(np.arange(0.0, 2.0 + 1e-9, 0.1), 4)


def header_row(ws, extra=None):
    ws.cell(row=1, column=1, value='时间\\到药材中心的距离')
    for j, rc in enumerate(R_CM):
        ws.cell(row=1, column=2 + j, value=round(float(rc), 4))
    if extra:
        ws.cell(row=1, column=2 + len(R_CM), value=extra)


def fill_sheet(ws, t_arr, field, n_dec=4, surf=None):
    for i, tt in enumerate(t_arr):
        ws.cell(row=2 + i, column=1, value=int(round(tt)))
        for j in range(len(R_CM)):
            v = field[i, j]
            if v is None or (isinstance(v, float) and np.isnan(v)):
                continue
            ws.cell(row=2 + i, column=2 + j, value=round(float(v), n_dec))
        if surf is not None:
            ws.cell(row=2 + i, column=2 + len(R_CM), value=round(float(surf[i]), n_dec))


def write_result1():
    d = json.load(open(os.path.join(_HERE, 'figures', 'problem_1_results.json'), encoding='utf-8'))
    T = np.array(d['T_field']); C = np.array(d['C_field'])
    t = np.arange(1, 1801)
    wb = Workbook()
    for name, field in (('温度', T), ('水分浓度', C)):
        ws = wb.active if name == '温度' else wb.create_sheet(name)
        ws.title = name
        header_row(ws)
        fill_sheet(ws, t, field)
    wb.save(os.path.join(OUTDIR, 'result1.xlsx'))
    print('[result1] %d行×21列 两表四位小数' % len(t))


def write_result2():
    # 直接使用 solve_q234.py 的 dt=1s 生产数据（problem_2_results.json，每隔1s真算）
    d = json.load(open(os.path.join(_HERE, 'figures', 'problem_2_results.json'), encoding='utf-8'))
    T = np.array(d['T_field']); C = np.array(d['C_field'])
    t = np.array(d['t_s'], dtype=int)
    wb = Workbook()
    for name, field in (('温度', T), ('水分浓度', C)):
        ws = wb.active if name == '温度' else wb.create_sheet(name)
        ws.title = name
        header_row(ws)
        fill_sheet(ws, t, field)
    wb.save(os.path.join(OUTDIR, 'result2.xlsx'))
    print('[result2] %d行×21列 两表四位小数 (dt=1s)' % len(t))


def _subsample60(t, *fields):
    """生产求解 dt=15s；xlsx 每 60s 一行，末行恒为烘干结束时刻 t*"""
    t = np.asarray(t, dtype=int)
    keep = (t % 60 == 0)
    keep[-1] = True  # 末行 t* 恒保留
    idx = np.where(keep)[0]
    return t[idx], [np.asarray(f)[idx] for f in fields]


def write_result3():
    d = json.load(open(os.path.join(_HERE, 'figures', 'problem_3_results.json'), encoding='utf-8'))
    C = np.array(d['C_field'])
    t = np.array(d['t_s'])
    t, (C,) = _subsample60(t, C)
    wb = Workbook()
    ws = wb.active
    header_row(ws)
    fill_sheet(ws, t, C)
    wb.save(os.path.join(OUTDIR, 'result3.xlsx'))
    print('[result3] %d行×21列 每60s至t*=%.0fs' % (len(t), d['meta']['t_star_s']))


def write_result4():
    d = json.load(open(os.path.join(_HERE, 'figures', 'problem_4_results.json'), encoding='utf-8'))
    C = np.array(d['C_field'], dtype=float)
    t = np.array(d['t_s'])
    surf = np.array(d['C_surf'], dtype=float)
    t, (C, surf) = _subsample60(t, C, surf)
    wb = Workbook()
    ws = wb.active
    header_row(ws, extra='药材表面')
    fill_sheet(ws, t, C, surf=surf)
    wb.save(os.path.join(OUTDIR, 'result4.xlsx'))
    print('[result4] %d行×21列+表面列 每60s至t*_shrink=%.0fs' % (len(t), d['meta']['t_star_s']))


if __name__ == '__main__':
    write_result1()
    write_result2()
    write_result3()
    write_result4()
    print('[fill_results] 附件3 四个文件已填充')

