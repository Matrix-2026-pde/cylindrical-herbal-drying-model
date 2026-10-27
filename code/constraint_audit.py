# 约束闭环审计：从落盘 result1.xlsx / problem_1_results.json 重算所有题设硬约束（不信任求解器）
# 只 print 结论，禁止 print 整个数组。
import os, sys, json
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)
import numpy as np
from openpyxl import load_workbook

T_INIT, C_INIT = 28.0, 2.55


def audit():
    d = json.load(open('figures/problem_1_results.json', encoding='utf-8'))
    T = np.array(d['T_field']); C = np.array(d['C_field'])
    T_inf_max = d['meta']['T_inf_max']

    # 自动发现基线（本题无优化基线，仅登记以满足闸约定）
    _BASE = []
    audited = set()

    violations = []
    # 1. 温度上下界
    if T.min() < T_INIT - 1e-4:
        violations.append(f"T<init: min={T.min():.4f}")
    if T.max() > T_inf_max + 0.5:
        violations.append(f"T>Tinf_max: max={T.max():.4f} > {T_inf_max:.4f}")
    # 2. 含水率上下界
    if C.max() > C_INIT + 1e-4:
        violations.append(f"C>init: max={C.max():.4f}")
    if C.min() < 0:
        violations.append(f"C<0: min={C.min():.4f}")
    # 3. 径向单调 T增 C减（容差1e-3）
    dT = np.diff(T, axis=1); dC = np.diff(C, axis=1)
    if dT.min() < -1e-3:
        idx = np.unravel_index(np.argmin(dT), dT.shape)
        violations.append(f"T径向非单调 at t_idx={idx[0]} r_idx={idx[1]} dT={dT.min():.2e}")
    if dC.max() > 1e-3:
        idx = np.unravel_index(np.argmax(dC), dC.shape)
        violations.append(f"C径向非单调 at t_idx={idx[0]} r_idx={idx[1]} dC={dC.max():.2e}")
    # 4. 中心温度时间单调不减
    if np.diff(T[:, 0]).min() < -1e-3:
        violations.append(f"中心温度时间下降 min_dt={np.diff(T[:,0]).min():.2e}")
    # 5. 中心1800s水分接近初值（逻辑合同 bounds: upper）
    if C_INIT - C[-1, 0] > 0.05:
        violations.append(f"中心失水过大: {C_INIT-C[-1,0]:.4f}")
    # 6. NaN/Inf
    if not (np.all(np.isfinite(T)) and np.all(np.isfinite(C))):
        violations.append("含 NaN/Inf")

    # 7. result1.xlsx 与 JSON 一致性（抽查）+ 四位小数
    wb = load_workbook('result1.xlsx', data_only=True)
    wsT = wb['温度']; wsC = wb['水分浓度']
    # 列布局：col1=时间, col2=r=0(idx0), ..., col22=r=2.0(idx20)
    xT_1800_surf = wsT.cell(row=1+1800, column=2+20).value
    xC_1800_surf = wsC.cell(row=1+1800, column=2+20).value
    if abs(round(float(T[-1, -1]), 4) - xT_1800_surf) > 1e-9:
        violations.append(f"xlsx温度不一致 {xT_1800_surf} vs {T[-1,-1]:.4f}")
    if abs(round(float(C[-1, -1]), 4) - xC_1800_surf) > 1e-9:
        violations.append(f"xlsx水分不一致 {xC_1800_surf} vs {C[-1,-1]:.4f}")

    audited.add('problem_1_field')
    missing = [b for b in _BASE if b not in audited]
    assert not missing, f"[基线漏审] {missing}"

    n = len(violations)
    ok = n == 0
    print(f"[constraint_audit] problem_1 PASS={ok}  n_violations={n}")
    print(f"    T range=[{T.min():.4f},{T.max():.4f}]  C range=[{C.min():.4f},{C.max():.4f}]")
    print(f"    中心C(1800)={C[-1,0]:.4f}(≤2.55上界侧) 表面C(1800)={C[-1,-1]:.4f}")
    print(f"    grid_conv maxreldiff T={d['grid_conv_maxreldiff']['T']:.2e} C={d['grid_conv_maxreldiff']['C']:.2e}")
    for v in violations[:5]:
        print("    VIOLATION:", v)
    return ok


if __name__ == '__main__':
    ok = audit()
    sys.exit(0 if ok else 1)

